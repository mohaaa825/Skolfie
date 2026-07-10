"""
file_parsers.py
----------------
Utilities for extracting raw text from uploaded files.

Supports:
    - .txt   (plain text)
    - .docx  (Word documents, via python-docx)
    - .pdf   (PDF documents, via PyPDF2)

Every parser function returns a single string of extracted text.
"""

import os


def read_txt(file_path: str) -> str:
    """Read a plain text file, trying a couple of common encodings."""
    for encoding in ("utf-8", "latin-1"):
        try:
            with open(file_path, "r", encoding=encoding) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    raise ValueError(f"Could not decode text file: {file_path}")


def read_docx(file_path: str) -> str:
    """Extract all paragraph (and table cell) text from a .docx file."""
    try:
        import docx
    except ImportError as exc:
        raise ImportError(
            "python-docx is required to read .docx files. "
            "Install it with: pip install python-docx"
        ) from exc

    document = docx.Document(file_path)
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

    # Reports and letters often store content in tables too.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)

    return "\n".join(paragraphs)


def read_pdf(file_path: str) -> str:
    """Extract text from a .pdf file, page by page."""
    try:
        import PyPDF2
    except ImportError as exc:
        raise ImportError(
            "PyPDF2 is required to read .pdf files. "
            "Install it with: pip install PyPDF2"
        ) from exc

    text_chunks = []
    with open(file_path, "rb") as f:
        reader = PyPDF2.PdfReader(f)
        for page in reader.pages:
            page_text = page.extract_text() or ""
            text_chunks.append(page_text)

    return "\n".join(text_chunks)


def extract_text_from_file(file_path: str) -> str:
    """
    Dispatch to the right parser based on file extension.
    Raises a ValueError for unsupported extensions.
    """
    if not file_path:
        return ""

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".txt":
        return read_txt(file_path)
    elif ext == ".docx":
        return read_docx(file_path)
    elif ext == ".pdf":
        return read_pdf(file_path)
    else:
        raise ValueError(
            f"Unsupported file type '{ext}'. Please upload a .txt, .docx, or .pdf file."
        )
