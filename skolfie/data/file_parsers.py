"""
skolfie/data/file_parsers.py
-----------------------------
Utilities for extracting raw text from uploaded files.

Supports:
    - .txt   (plain text)
    - .docx  (Word documents, via python-docx)
    - .pdf   (PDF documents, via pypdf)

Every parser function returns a single string of extracted text.

Uploads are untrusted, so the parsers bound the work they do: the file size
(`MAX_FILE_BYTES`), the uncompressed size of a .docx zip
(`MAX_DOCX_UNCOMPRESSED_BYTES`, against zip bombs) and the number of PDF
pages (`MAX_PDF_PAGES`). The content must match the extension (a .docx must
be a Word zip, a .pdf must start with a PDF header, a .txt must not be
binary). Any problem raises `FileParseError`, whose message is safe to show
to users (it never contains the server-side path).
"""

from __future__ import annotations

import codecs
import os
import zipfile

SUPPORTED_EXTENSIONS = (".txt", ".docx", ".pdf")

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_DOCX_MEMBERS = 5000
MAX_PDF_PAGES = 500


class FileParseError(ValueError):
    """An upload could not be read; the message is meant for the user."""


def read_txt(file_path: str) -> str:
    """
    Read a plain text file.

    A byte-order mark (UTF-8 / UTF-16) wins if present; otherwise UTF-8 is
    tried, then Windows-1252, then Latin-1 (which accepts any byte sequence).
    Text containing NUL characters is rejected as binary.
    """
    with open(file_path, "rb") as f:
        raw = f.read()

    try:
        if raw.startswith(codecs.BOM_UTF8):
            text = raw.decode("utf-8-sig")
        elif raw.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
            text = raw.decode("utf-16")
        else:
            text = None
            for encoding in ("utf-8", "cp1252"):
                try:
                    text = raw.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            if text is None:
                text = raw.decode("latin-1")
    except UnicodeDecodeError as exc:
        raise FileParseError("This .txt file has a broken text encoding.") from exc

    if "\x00" in text:
        raise FileParseError("This .txt file looks like a binary file, not text.")
    return text


def _check_docx_zip(file_path: str, max_uncompressed: int) -> None:
    """Reject non-zip files and zip bombs before python-docx unpacks anything."""
    try:
        with zipfile.ZipFile(file_path) as zf:
            members = zf.infolist()
    except (zipfile.BadZipFile, OSError) as exc:
        raise FileParseError("This is not a valid .docx file.") from exc
    if not any(m.filename == "word/document.xml" for m in members):
        raise FileParseError("This is not a valid .docx file.")
    if len(members) > MAX_DOCX_MEMBERS:
        raise FileParseError("This .docx file has too many parts to process.")
    # zipfile never returns more than an entry's declared file_size, so the
    # declared sizes bound what python-docx can decompress.
    if sum(m.file_size for m in members) > max_uncompressed:
        raise FileParseError(
            f"This .docx file is too large once unpacked (limit "
            f"{max_uncompressed // (1024 * 1024)} MB)."
        )


def read_docx(file_path: str, max_uncompressed: int = MAX_DOCX_UNCOMPRESSED_BYTES) -> str:
    """Extract all paragraph (and table cell) text from a .docx file."""
    try:
        import docx
    except ImportError as exc:
        raise ImportError(
            "python-docx is required to read .docx files. "
            "Install it with: pip install python-docx"
        ) from exc

    _check_docx_zip(file_path, max_uncompressed)
    try:
        document = docx.Document(file_path)
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

        # Reports and letters often store content in tables too.
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text)
    except Exception as exc:  # python-docx / lxml / zipfile errors on corrupt files
        raise FileParseError("Could not read this .docx file; it may be corrupt.") from exc

    return "\n".join(paragraphs)


def read_pdf(file_path: str, max_pages: int = MAX_PDF_PAGES) -> str:
    """Extract text from a .pdf file, page by page (at most `max_pages` pages)."""
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ImportError(
            "pypdf is required to read .pdf files. "
            "Install it with: pip install pypdf"
        ) from exc

    with open(file_path, "rb") as f:
        if b"%PDF-" not in f.read(1024):
            raise FileParseError("This is not a valid .pdf file.")
        f.seek(0)
        try:
            reader = PdfReader(f)
            if reader.is_encrypted:
                # pypdf already tried the empty password; anything else is locked.
                try:
                    reader.decrypt("")
                    reader.pages[0]
                except Exception as exc:
                    raise FileParseError(
                        "This PDF is password-protected. Remove the password and try again."
                    ) from exc
            if len(reader.pages) > max_pages:
                raise FileParseError(
                    f"This PDF has {len(reader.pages)} pages; the limit is {max_pages}."
                )
            text_chunks = [page.extract_text() or "" for page in reader.pages]
        except FileParseError:
            raise
        except Exception as exc:  # pypdf errors on corrupt or hostile files
            raise FileParseError("Could not read this .pdf file; it may be corrupt.") from exc

    return "\n".join(text_chunks)


def extract_text_from_file(file_path: str | os.PathLike[str],
                           max_bytes: int = MAX_FILE_BYTES) -> str:
    """
    Dispatch to the right parser based on file extension.

    Raises FileParseError (a ValueError) for unsupported extensions, files
    over `max_bytes`, and files that are corrupt or don't match their extension.
    """
    if not file_path:
        return ""

    file_path = os.fspath(file_path)
    ext = os.path.splitext(file_path)[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise FileParseError(
            f"Unsupported file type '{ext}'. Please upload a .txt, .docx, or .pdf file."
        )

    try:
        size = os.path.getsize(file_path)
    except OSError as exc:
        raise FileParseError("The uploaded file could not be opened.") from exc
    if size > max_bytes:
        raise FileParseError(
            f"The file is too large ({size / (1024 * 1024):.1f} MB); "
            f"the limit is {max_bytes / (1024 * 1024):.0f} MB."
        )

    if ext == ".txt":
        return read_txt(file_path)
    if ext == ".docx":
        return read_docx(file_path)
    return read_pdf(file_path)
